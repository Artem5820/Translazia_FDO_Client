from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parent
REPORT = ROOT / "report.md"
OUT = ROOT / "lab8.docx"


def style_run(run, size=14, font="Times New Roman"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = False
    run.italic = False


def add_paragraph(doc: Document, text: str = "", *, style: str | None = None, first_line=True):
    paragraph = doc.add_paragraph(style=style)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(1.25 if first_line else 0)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    style_run(run)
    return paragraph


def add_code(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    style_run(run, size=10, font="Courier New")


def add_table_from_lines(doc: Document, lines: list[str]):
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        rows.append(cells)
    if len(rows) < 2:
        return
    rows = [rows[0]] + rows[2:]
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    for idx, value in enumerate(rows[0]):
        table.rows[0].cells[idx].text = value
    for row_values in rows[1:]:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            row.cells[idx].text = value
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    style_run(run, size=11)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)

    in_code = False
    table_buffer: list[str] = []
    for raw_line in REPORT.read_text(encoding="utf-8").splitlines():
        line = raw_line.rstrip()
        if line.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            add_code(doc, line)
            continue
        if line.startswith("|"):
            table_buffer.append(line)
            continue
        if table_buffer:
            add_table_from_lines(doc, table_buffer)
            table_buffer = []
        if not line:
            add_paragraph(doc, "")
        elif line.startswith("# "):
            paragraph = add_paragraph(doc, line[2:], first_line=False)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            add_paragraph(doc, line[3:], first_line=False)
        elif line.startswith("### "):
            add_paragraph(doc, line[4:], first_line=True)
        else:
            add_paragraph(doc, line)
    if table_buffer:
        add_table_from_lines(doc, table_buffer)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
