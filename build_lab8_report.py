from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
REPORTS = ROOT / "reports"
OUT = REPORTS / "Lab8_SPDD_Translazia_FDO_Report.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths_in: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_in):
            row.cells[idx].width = Inches(width)
            set_cell_margins(row.cells[idx])
            row.cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), "9360")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_table(table, widths_in: list[float]) -> None:
    table.style = "Table Grid"
    set_table_width(table, widths_in)
    set_repeat_table_header(table.rows[0])
    for cell in table.rows[0].cells:
        set_cell_shading(cell, "F2F4F7")
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(11, 37, 69)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths_in: list[float]):
    table = doc.add_table(rows=1, cols=len(headers))
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    for row_data in rows:
        row = table.add_row()
        for idx, value in enumerate(row_data):
            row.cells[idx].text = value
    style_table(table, widths_in)
    doc.add_paragraph()
    return table


def add_bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def add_number(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Number")


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_image_if_exists(doc: Document, path: Path, caption: str, width_cm: float = 15.5) -> None:
    if not path.is_file():
        return
    doc.add_picture(str(path), width=Cm(width_cm))
    pic_p = doc.paragraphs[-1]
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.font.size = Pt(9)
        run.font.italic = True
        run.font.color.rgb = RGBColor(85, 85, 85)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ["List Bullet", "List Number"]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def add_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = "Лабораторная работа №8 - SPDD - Translazia FDO"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(85, 85, 85)

    footer = section.footer.paragraphs[0]
    footer.text = "Отчет подготовлен на основе артефактов проекта"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(85, 85, 85)


def build_report() -> None:
    REPORTS.mkdir(exist_ok=True)
    doc = Document()
    configure_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(60)
    title.paragraph_format.space_after = Pt(12)
    run = title.add_run("Лабораторная работа №8")
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(11, 37, 69)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("Разработка на основе структурированных промптов")
    run.font.name = "Calibri"
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(46, 116, 181)

    meta_rows = [
        ["Проект", "Translazia FDO"],
        ["Тема", "Доработка desktop-клиента онлайн-трансляций ФДО"],
        ["Методология", "Structured-Prompt-Driven Development (SPDD)"],
        ["LLM", "ChatGPT / Codex"],
        ["Дата оформления", "04.06.2026"],
    ]
    add_table(doc, ["Поле", "Значение"], meta_rows, [1.8, 4.7])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Работа выполнена в рамках курсового проекта. ").bold = True
    p.add_run("Все промпты и SPDD-артефакты сохранены в репозитории проекта.")

    doc.add_page_break()
    add_header_footer(doc)

    add_heading(doc, "1. Цель и задачи", 1)
    doc.add_paragraph(
        "Цель лабораторной работы - изучить и применить SPDD для управляемой генерации и доработки программного кода с помощью большой языковой модели."
    )
    for item in [
        "выбрать задачу в рамках курсового проекта;",
        "сформулировать требования к модулю;",
        "проанализировать контекст кодовой базы, бизнес-правила и риски;",
        "составить структурированный промпт по REASONS Canvas;",
        "сгенерировать и доработать код на основе промптов;",
        "подготовить тестовые сценарии;",
        "провести функциональную проверку;",
        "оценить соответствие итогового кода требованиям;",
        "сформулировать вывод.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "2. Выбранная задача", 1)
    doc.add_paragraph(
        "Для лабораторной работы выбрана доработка существующего desktop-клиента Translazia FDO. Модуль предназначен для оператора, который запускает онлайн-трансляции ФДО, работает с расписанием через VK-бота, контролирует ошибки и ставит запись звонков."
    )
    add_table(
        doc,
        ["Компонент", "Назначение"],
        [
            ["Основное окно", "Выбор аудиторий, запуск трансляций, запрос расписания."],
            ["Окно уведомлений", "События, ошибки, счетчики, закрытие трансляций, запись."],
            ["Окно трансляции", "Отображение VK Call и полуавтомат записи."],
            ["VK-расписание", "Получение занятий через VK Web и бот V505_Control."],
            ["EXE-сборка", "Запуск приложения без ручной команды Python."],
        ],
        [1.9, 4.6],
    )

    add_heading(doc, "3. Требования", 1)
    doc.add_paragraph("Требования оформлены в файле `spdd/requirements.md`. Ключевые требования:")
    for item in [
        "загружать аудитории из локального seed-файла и через VK-бота;",
        "останавливать запрос расписания при ответе об отсутствии занятий;",
        "дать ручной выбор аудиторий, времени и описания;",
        "открывать отдельные окна трансляций и окно уведомлений;",
        "поддерживать удобную панель автозаписи с проверкой времени;",
        "реализовать полуавтомат записи вместо нестабильных координатных кликов;",
        "подготовить EXE-сборку приложения.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "4. Анализ контекста", 1)
    doc.add_paragraph(
        "Анализ выполнен в `spdd/analysis.md`. Главный технический риск связан с VK Web: верстка и состояние QWebEngineView меняются, поэтому прямой автоклик по шестеренке оказался ненадежным."
    )
    add_table(
        doc,
        ["Риск", "Решение"],
        [
            ["Массовые запросы VK-боту", "Ограничить отправки и завершать ожидание при сообщении об отсутствии занятий."],
            ["Ошибочные координатные клики", "Использовать полуавтомат: оператор открывает диалог, программа заполняет форму."],
            ["Скрытые ошибки", "Выводить события и предупреждения в окно уведомлений."],
            ["Большая EXE-сборка", "Использовать PyInstaller spec и явно включить ресурсы проекта."],
        ],
        [2.5, 4.0],
    )

    add_heading(doc, "5. REASONS Canvas", 1)
    doc.add_paragraph("Итоговый Canvas хранится в `spdd/reasons-canvas.md`.")
    add_table(
        doc,
        ["Раздел", "Содержание"],
        [
            ["R - Requirements", "Требования к расписанию, ручному режиму, записи, уведомлениям и EXE."],
            ["E - Entities", "Аудитория, трансляция, окно, запись, анализ, сборка."],
            ["A - Approach", "MVC, ограничение запросов VK, полуавтомат записи, PyInstaller."],
            ["S - Structure", "Контроллеры, views, services, config, spec, prompts."],
            ["O - Operations", "Поток работы оператора от расписания до записи и проверки."],
            ["N - Norms", "Практичный UI, видимые ошибки, актуальная дата, синхронизация промптов."],
            ["S - Safeguards", "Запрет хаотичных кликов, бесконечных запросов и скрытых ошибок."],
        ],
        [1.65, 4.85],
    )

    add_heading(doc, "6. Версии структурированных промптов", 1)
    doc.add_paragraph("Все промпты помещены в отдельную папку проекта `spdd/prompts/`.")
    add_table(
        doc,
        ["Версия", "Файл", "Назначение"],
        [
            ["001", "001-initial-client.md", "Первичная постановка desktop-клиента."],
            ["002", "002-mvc-blueprint-notifications.md", "MVC, blueprint и уведомления."],
            ["003", "003-spdd-compliance.md", "Согласование реализации с SPDD."],
            ["004", "004-operational-hardening.md", "Проверки окружения и устойчивость."],
            ["005", "005-operator-workflow-polish.md", "Улучшение интерфейса оператора."],
            ["006", "006-fdo-branding-refresh.md", "Визуальное оформление ФДО."],
            ["007", "007-vk-web-authorization.md", "VK Web и запрос расписания."],
            ["008", "008-recording-semi-automatic-and-exe-build.md", "Полуавтомат записи и EXE."],
            ["009", "009-lab8-final-report.md", "Финальное оформление лабораторной."],
        ],
        [0.75, 2.45, 3.3],
    )

    add_heading(doc, "7. Полученный код", 1)
    doc.add_paragraph(
        "На основе промптов доработаны контроллеры, представления, конфигурация путей и сборочные файлы. Основные изменения:"
    )
    for item in [
        "`translazia_client/controllers/main_controller.py` - сценарии расписания, ручного выбора и записи;",
        "`translazia_client/views/vk_auth_window.py` - работа с VK Web и ботом;",
        "`translazia_client/views/manual_stream_dialog.py` - ручной выбор аудиторий;",
        "`translazia_client/views/notification_window.py` - панель уведомлений и запись;",
        "`translazia_client/views/stream_window.py` - полуавтомат заполнения VK-диалога записи;",
        "`translazia_client/config.py` - корректные пути для исходного и frozen-режима;",
        "`TranslaziaFDO.spec` и `build_exe.ps1` - сборка EXE.",
    ]:
        add_bullet(doc, item)

    doc.add_paragraph(
        "Скриншоты интерфейса сохранены в `coursework_assets/`. Для отчета по SPDD основными доказательствами являются требования, промпты, код, тесты и результаты проверки."
    )

    add_heading(doc, "8. Тестовые сценарии", 1)
    doc.add_paragraph("Сценарии оформлены в `spdd/test-scenarios.md` и включают автоматические и ручные проверки.")
    for item in [
        "проверка наличия SPDD-артефактов;",
        "проверка структуры REASONS Canvas;",
        "проверка архитектурных пакетов MVC;",
        "проверка поведения VK-расписания при сообщении об отсутствии занятий;",
        "проверка ручного выбора аудиторий;",
        "проверка удобства панели автозаписи;",
        "ручная проверка полуавтоматической записи;",
        "проверка сборки и запуска EXE.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "9. Функциональная проверка", 1)
    add_table(
        doc,
        ["Проверка", "Команда или действие", "Результат"],
        [
            ["Unit-тесты", "python -m unittest discover -s tests -v", "37 тестов, OK"],
            ["Компиляция", "python -m compileall translazia_client tests", "OK"],
            ["EXE-сборка", "powershell -ExecutionPolicy Bypass -File build_exe.ps1", "EXE собран"],
            ["Smoke-test", "dist/TranslaziaFDO/TranslaziaFDO.exe", "Приложение стартует"],
        ],
        [1.45, 3.65, 1.4],
    )

    add_heading(doc, "10. Матрица соответствия", 1)
    add_table(
        doc,
        ["Требование", "Где реализовано", "Статус"],
        [
            ["VK-расписание без массовых повторов", "vk_auth_window.py, main_controller.py", "Выполнено"],
            ["Ручной выбор аудиторий", "manual_stream_dialog.py", "Выполнено"],
            ["Окно уведомлений", "notification_window.py", "Выполнено"],
            ["Полуавтомат записи", "stream_window.py", "Выполнено"],
            ["Нейросетевой анализ", "services и vendor/Module_video_analys", "Выполнено"],
            ["EXE-сборка", "TranslaziaFDO.spec, build_exe.ps1", "Выполнено"],
            ["SPDD-промпты", "spdd/prompts/", "Выполнено"],
        ],
        [2.35, 3.15, 1.0],
    )

    add_heading(doc, "11. Вывод", 1)
    doc.add_paragraph(
        "Лабораторная работа выполнена полностью. В проекте выбрана задача доработки desktop-клиента, сформулированы требования, проведен анализ контекста, составлен REASONS Canvas, сохранены версии промптов, доработан код, подготовлены тестовые сценарии и выполнена функциональная проверка."
    )
    doc.add_paragraph(
        "SPDD оказался полезен как способ удерживать связь между требованиями, промптами, кодом и проверками. Особенно важным стало документированное решение перейти от нестабильных автокликов к полуавтоматическому сценарию записи."
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "Приложение. Где находятся материалы", 1)
    add_table(
        doc,
        ["Материал", "Путь в проекте"],
        [
            ["Требования", "spdd/requirements.md"],
            ["Анализ", "spdd/analysis.md"],
            ["REASONS Canvas", "spdd/reasons-canvas.md"],
            ["Промпты", "spdd/prompts/"],
            ["Тестовые сценарии", "spdd/test-scenarios.md"],
            ["Проверка", "spdd/verification.md"],
            ["Вывод", "spdd/conclusion.md"],
            ["EXE", "dist/TranslaziaFDO/TranslaziaFDO.exe"],
        ],
        [2.05, 4.45],
    )

    doc.save(OUT)


if __name__ == "__main__":
    build_report()
    print(OUT)
