from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


DOCX = "Пояснительная записка SoundChecker.docx"


def math_run(text: str):
    mr = OxmlElement("m:r")
    mt = OxmlElement("m:t")
    mt.text = text
    mr.append(mt)
    return mr


def math_text(text: str):
    omath = OxmlElement("m:oMath")
    omath.append(math_run(text))
    return omath


def math_stft():
    omath = OxmlElement("m:oMath")
    omath.append(math_run("X(k,m)="))

    nary = OxmlElement("m:nary")
    nary_pr = OxmlElement("m:naryPr")
    chr_el = OxmlElement("m:chr")
    chr_el.set(qn("m:val"), "∑")
    lim_loc = OxmlElement("m:limLoc")
    lim_loc.set(qn("m:val"), "undOvr")
    nary_pr.append(chr_el)
    nary_pr.append(lim_loc)
    nary.append(nary_pr)

    sub = OxmlElement("m:sub")
    sub.append(math_run("n=0"))
    sup = OxmlElement("m:sup")
    sup.append(math_run("N−1"))
    expr = OxmlElement("m:e")
    expr.append(math_run("x(n)·w(n−mH)·"))

    exp = OxmlElement("m:sSup")
    base = OxmlElement("m:e")
    base.append(math_run("e"))
    power = OxmlElement("m:sup")
    power.append(math_run("−j2πkn/N"))
    exp.append(base)
    exp.append(power)
    expr.append(exp)

    nary.append(sub)
    nary.append(sup)
    nary.append(expr)
    omath.append(nary)
    return omath


def set_equation_paragraph(paragraph, equation_element, number: str) -> None:
    paragraph._p.clear_content()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph._p.append(equation_element)
    run = paragraph.add_run(f"        ({number})")
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(14)
    run.bold = False
    run.italic = False


def main() -> None:
    doc = Document(DOCX)
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("L = 20 lg(A)"):
            set_equation_paragraph(paragraph, math_text("L = 20·lg(A)"), "1")
        elif text.startswith("X(k, m) = сумма"):
            set_equation_paragraph(paragraph, math_stft(), "2")
    doc.save(DOCX)


if __name__ == "__main__":
    main()
